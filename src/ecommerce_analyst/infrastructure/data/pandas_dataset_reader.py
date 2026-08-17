"""Pandas-backed dataset reader for CSV and XLSX uploads."""

from __future__ import annotations

import csv
from collections.abc import Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import pandas as pd
from openpyxl import load_workbook

from ecommerce_analyst.domain.enums import DatasetFileType
from ecommerce_analyst.domain.exceptions import DatasetIngestionError
from ecommerce_analyst.domain.models import (
    CellValue,
    ColumnMetadata,
    DatasetIngestionLimits,
    DatasetMetadata,
    IngestedDataset,
)


class PandasDatasetReader:
    """Read supported uploaded datasets into a domain-neutral representation."""

    def read(self, file_path: Path, limits: DatasetIngestionLimits) -> IngestedDataset:
        """Read a CSV or XLSX file with technical validation only."""

        resolved_path = Path(file_path)
        file_type = _resolve_file_type(resolved_path, limits)
        file_size_bytes = _validate_file_size(resolved_path, limits)

        match file_type:
            case DatasetFileType.CSV:
                header = _read_csv_header(resolved_path)
                _validate_header(header)
                dataframe = _read_csv(resolved_path)
            case DatasetFileType.XLSX:
                header = _read_xlsx_header(resolved_path)
                _validate_header(header)
                dataframe = _read_xlsx(resolved_path)

        _validate_loaded_dataframe(dataframe)
        dataframe.columns = header

        warnings = _header_warnings(header)
        metadata = DatasetMetadata(
            filename=resolved_path.name,
            file_type=file_type,
            file_size_bytes=file_size_bytes,
            row_count=len(dataframe.index),
            column_count=len(dataframe.columns),
            columns=tuple(
                ColumnMetadata(index=index, name=str(name), inferred_dtype=str(dtype))
                for index, (name, dtype) in enumerate(dataframe.dtypes.items())
            ),
            parsing_warnings=tuple(warnings),
        )

        return IngestedDataset(metadata=metadata, records=_records_from_dataframe(dataframe))


def _resolve_file_type(file_path: Path, limits: DatasetIngestionLimits) -> DatasetFileType:
    if not file_path.exists():
        raise DatasetIngestionError(
            f"Dataset file does not exist: {file_path}", code="file_not_found"
        )
    if not file_path.is_file():
        raise DatasetIngestionError(f"Dataset path is not a file: {file_path}", code="not_a_file")

    try:
        file_type = DatasetFileType.from_extension(file_path.suffix)
    except ValueError as exc:
        raise DatasetIngestionError(
            "Unsupported dataset format. Upload a CSV or XLSX file.",
            code="unsupported_file_type",
        ) from exc

    if file_type not in limits.supported_file_types:
        raise DatasetIngestionError(
            f"{file_type.value.upper()} files are not enabled for ingestion.",
            code="file_type_disabled",
        )

    return file_type


def _validate_file_size(file_path: Path, limits: DatasetIngestionLimits) -> int:
    file_size_bytes = file_path.stat().st_size
    if file_size_bytes > limits.max_file_size_bytes:
        raise DatasetIngestionError(
            (
                f"Dataset file is {file_size_bytes} bytes, which exceeds the configured "
                f"limit of {limits.max_file_size_bytes} bytes."
            ),
            code="file_too_large",
        )
    return file_size_bytes


def _read_csv_header(file_path: Path) -> list[str]:
    try:
        with file_path.open("r", encoding="utf-8-sig", newline="") as csv_file:
            reader = csv.reader(csv_file)
            raw_header = next(reader, None)
    except UnicodeDecodeError as exc:
        raise DatasetIngestionError(
            "CSV file could not be decoded as UTF-8.",
            code="encoding_error",
        ) from exc
    except csv.Error as exc:
        raise DatasetIngestionError(
            "CSV header could not be parsed.",
            code="parse_error",
        ) from exc

    if raw_header is None:
        raise DatasetIngestionError("Dataset file is empty.", code="empty_dataset")
    return [str(column) for column in raw_header]


def _read_xlsx_header(file_path: Path) -> list[str]:
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        worksheet = workbook.active
        if worksheet is None:
            workbook.close()
            raise DatasetIngestionError(
                "XLSX workbook contains no active sheet.", code="empty_dataset"
            )
        raw_header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        workbook.close()
    except (BadZipFile, ValueError, OSError) as exc:
        raise DatasetIngestionError(
            "XLSX file could not be opened as a valid Excel workbook.",
            code="parse_error",
        ) from exc

    if raw_header is None:
        raise DatasetIngestionError("Dataset file is empty.", code="empty_dataset")
    return ["" if column is None else str(column) for column in raw_header]


def _validate_header(header: Sequence[str]) -> None:
    if not header:
        raise DatasetIngestionError("Dataset is missing a header row.", code="missing_header")

    missing_indexes = [index for index, name in enumerate(header, start=1) if not name.strip()]
    if missing_indexes:
        joined_indexes = ", ".join(str(index) for index in missing_indexes)
        raise DatasetIngestionError(
            f"Dataset header contains blank column names at positions: {joined_indexes}.",
            code="missing_header",
        )

    seen: set[str] = set()
    duplicates: set[str] = set()
    for name in header:
        if name in seen:
            duplicates.add(name)
        seen.add(name)

    if duplicates:
        duplicate_names = ", ".join(sorted(duplicates))
        raise DatasetIngestionError(
            f"Dataset header contains duplicate column names: {duplicate_names}.",
            code="duplicate_columns",
        )


def _read_csv(file_path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(file_path, encoding="utf-8-sig", on_bad_lines="error")
    except UnicodeDecodeError as exc:
        raise DatasetIngestionError(
            "CSV file could not be decoded as UTF-8.",
            code="encoding_error",
        ) from exc
    except (pd.errors.EmptyDataError, pd.errors.ParserError, ValueError) as exc:
        raise DatasetIngestionError(
            "CSV file could not be parsed as a valid tabular dataset.",
            code="parse_error",
        ) from exc


def _read_xlsx(file_path: Path) -> pd.DataFrame:
    try:
        return pd.read_excel(file_path, engine="openpyxl")
    except (BadZipFile, ValueError, OSError) as exc:
        raise DatasetIngestionError(
            "XLSX file could not be parsed as a valid tabular dataset.",
            code="parse_error",
        ) from exc


def _validate_loaded_dataframe(dataframe: pd.DataFrame) -> None:
    if dataframe.shape[1] == 0:
        raise DatasetIngestionError("Dataset is missing a header row.", code="missing_header")
    if dataframe.shape[0] == 0:
        raise DatasetIngestionError("Dataset contains no data rows.", code="empty_dataset")


def _header_warnings(header: Sequence[str]) -> list[str]:
    warnings: list[str] = []
    whitespace_columns = [name for name in header if name != name.strip()]
    if whitespace_columns:
        warnings.append("One or more column names contain leading or trailing whitespace.")
    return warnings


def _records_from_dataframe(dataframe: pd.DataFrame) -> list[dict[str, CellValue]]:
    records: list[dict[str, CellValue]] = []
    object_dataframe = dataframe.astype(object).where(pd.notna(dataframe), None)

    for raw_record in object_dataframe.to_dict(orient="records"):
        record = {str(column): _to_cell_value(value) for column, value in raw_record.items()}
        records.append(record)

    return records


def _to_cell_value(value: Any) -> CellValue:
    if value is None or isinstance(value, str | int | float | bool | date | datetime):
        return value
    if hasattr(value, "item"):
        scalar_value = value.item()
        if isinstance(scalar_value, str | int | float | bool | date | datetime):
            return scalar_value
    return str(value)
